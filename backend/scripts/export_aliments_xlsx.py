"""Exporte le catalogue nutrition `aliments.csv` (transposé, `;`) en Excel.

Deux feuilles, même layout transposé que le CSV (ligne 1 = noms d'aliments,
colonne 1 = propriétés) :
  - "Aliments (Super C)"       : le catalogue courant (aliments.csv).
  - "Non achetables Super C"   : les items retirés du catalogue le 2026-07-17
    (poudres/barres protéinées Inshape/ON, whey Leanfit Costco — Super C ne les
    vend pas), lus depuis le backup pré-bascule pour ne rien perdre. Vide si
    le backup est absent.

Usage (depuis backend/) :
    uv run python -m scripts.export_aliments_xlsx [--out PATH]
"""
from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from app.core.config import settings

_CATALOG_DIR = settings.imports_dir / "Sante" / "tableur"
_DEFAULT_CSV = _CATALOG_DIR / "aliments.csv"
_BACKUP_CSV = _CATALOG_DIR / "aliments.pre-superc-2026-07-17.csv.bak"
_RETIRED_MARKERS = ("(inshape)", "(on)", "(costco)")


def _read_csv(path: Path) -> list[list[str]]:
    with path.open(encoding="utf-8", newline="") as f:
        return [row for row in csv.reader(f, delimiter=";")]


def _cell_value(raw: str) -> float | str | None:
    """Convertit en float si possible (virgule ou point décimal), sinon texte tel quel."""
    if raw == "":
        return None
    try:
        return float(raw.replace(",", "."))
    except ValueError:
        return raw


def _write_sheet(ws: Worksheet, rows: list[list[str]]) -> None:
    for row in rows:
        ws.append([row[0]] + [_cell_value(v) for v in row[1:]])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for cell in ws["A"]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 22


def _retired_columns(backup_rows: list[list[str]]) -> list[list[str]]:
    """Sous-table du backup limitée aux colonnes marquées (Inshape)/(ON)/(Costco)."""
    if not backup_rows:
        return []
    header = backup_rows[0]
    keep = [0] + [
        i for i, name in enumerate(header)
        if i > 0 and any(m in name.lower() for m in _RETIRED_MARKERS)
    ]
    return [[row[i] if i < len(row) else "" for i in keep] for row in backup_rows]


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--csv", type=Path, default=_DEFAULT_CSV)
    ap.add_argument("--backup", type=Path, default=_BACKUP_CSV)
    ap.add_argument("--out", type=Path, default=_CATALOG_DIR / "aliments.xlsx")
    args = ap.parse_args()

    if not args.csv.exists():
        print(f"[export] catalogue introuvable : {args.csv}")
        return 1

    wb = Workbook()
    ws = wb.active
    ws.title = "Aliments (Super C)"
    _write_sheet(ws, _read_csv(args.csv))

    ws2 = wb.create_sheet("Non achetables Super C")
    n_retired = 0
    if args.backup.exists():
        retired = _retired_columns(_read_csv(args.backup))
        n_retired = len(retired[0]) - 1 if retired else 0
        if n_retired > 0:
            _write_sheet(ws2, retired)
        else:
            ws2.append(["Aucun item retiré trouvé dans le backup."])
    else:
        ws2.append([f"Backup introuvable ({args.backup.name}) : feuille vide."])

    wb.save(args.out)
    n_foods = len(_read_csv(args.csv)[0]) - 1
    print(f"[export] {args.out}")
    print(f"[export]   Aliments (Super C) : {n_foods} aliments")
    print(f"[export]   Non achetables Super C : {n_retired} items")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

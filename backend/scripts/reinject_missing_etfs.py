"""Réinjecte les ETF présents dans ETF_Defensif (ToutBroker.xlsx) mais absents
de Sheet1 -- trouvés lors d'une vérification (~3241 ETF importés lors d'un
lot USA, jamais ajoutés à Sheet1). Sans une ligne Sheet1 marquée
Secteur 1 = "ETF", `_check_is_etf()` ne les reconnaît pas : le runner tente de
télécharger leurs comptes financiers (qu'un ETF n'a pas), échoue, et les
SUPPRIME de tickers.csv comme "faux délistés" -- cf. le run qui a fait
chuter tickers.csv de 10 472 à 7 077 lignes.

Étapes :
  1. tickers.csv (headerless, sep=";") : ajoute les tickers ETF_Defensif absents.
  2. Sheet1 de ToutBroker.xlsx : ajoute une ligne par ticker manquant, avec
     Secteur="ETF", Secteur 1="ETF", Secteur 2-5 classés par mots-clés du nom
     (réutilise `classify()` de classify_etf_sectors.py). Le reste (Prix,
     Volume, Achat, Chance MOAT...) sera rempli par le prochain run Buffett
     (chemin ETF normal, Score=200 fixe).

⚠️ Écrit Sheet1 via openpyxl EN PLACE (pas pandas `to_excel`, qui réécrirait
le classeur entier et effacerait ETF_Defensif/ETF_Pays -- déjà 2 sheets
supplémentaires sur ce fichier). Backup horodaté des deux fichiers avant
toute écriture.

Usage (depuis `backend/`) : `uv run python -m scripts.reinject_missing_etfs [--dry]`
"""
from __future__ import annotations

import csv
import datetime as dt
import sys
from pathlib import Path

import openpyxl

_BACKEND = Path(__file__).resolve().parents[1]
if str(_BACKEND) not in sys.path:
    sys.path.insert(0, str(_BACKEND))

from app.services.backup_storage import backup_file  # noqa: E402
from app.services.finance.buffett.broker_availability import find_broker_file  # noqa: E402
from app.services.finance.buffett.config import Config  # noqa: E402
from scripts.classify_etf_sectors import classify  # noqa: E402

SHEET1_HEADER_ROW = 1


_ETFUSA_CSV = (
    Path(__file__).resolve().parents[2]
    / "data" / "imports" / "Finances" / "variables" / "Ticker import" / "ETFUSA.csv"
)


def _load_etfusa_names(path: Path = _ETFUSA_CSV) -> dict[str, str]:
    """{ticker: nom réel} depuis la source d'import USA d'origine -- ETF_Defensif
    contient le TICKER comme "Nom" pour 3246/3488 lignes (jamais réellement
    rempli lors de l'enrichissement), ce qui réduirait `classify()` à un
    fourre-tout générique faute de mots-clés reconnaissables."""
    if not path.exists():
        return {}
    out: dict[str, str] = {}
    with open(path, encoding="utf-8", newline="") as f:
        for row in csv.DictReader(f):
            sym, name = row.get("SYMBOL"), row.get("NAME")
            if sym and name:
                out[sym.strip().upper()] = name.strip()
    return out


def _load_etf_defensif(path: Path) -> dict[str, str]:
    """{ticker: nom} depuis la feuille ETF_Defensif, avec repli sur le nom réel
    de `ETFUSA.csv` quand ETF_Defensif n'a que le ticker en guise de nom."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["ETF_Defensif"]
    rows = ws.iter_rows(values_only=True)
    next(rows, None)  # en-têtes : Ticker, Nom, Defensif_pct
    usa_names = _load_etfusa_names()
    out: dict[str, str] = {}
    for row in rows:
        if not row or not row[0]:
            continue
        t = str(row[0]).strip().upper()
        nom = str(row[1]).strip() if row[1] else t
        if nom.upper() == t and t in usa_names:
            nom = usa_names[t]
        out[t] = nom
    wb.close()
    return out


def _load_tickers_csv(path: Path) -> set[str]:
    if not path.exists():
        return set()
    out: set[str] = set()
    with open(path, encoding="utf-8", newline="") as f:
        for row in csv.reader(f, delimiter=";"):
            if row and row[0].strip():
                out.add(row[0].strip().upper())
    return out


def main(write: bool = True) -> None:
    bf = Path(find_broker_file())
    tk_path = Path(Config.TICKERS_CSV)

    etf_defensif = _load_etf_defensif(bf)
    print(f"ETF_Defensif : {len(etf_defensif)} tickers")

    wb = openpyxl.load_workbook(bf)
    ws = wb["Sheet1"]
    header = [c.value for c in ws[SHEET1_HEADER_ROW]]
    col = {name: i + 1 for i, name in enumerate(header)}  # 1-indexé (openpyxl)

    sheet1_tickers: set[str] = set()
    for row in ws.iter_rows(min_row=SHEET1_HEADER_ROW + 1, values_only=True):
        if row and row[0]:
            sheet1_tickers.add(str(row[0]).strip().upper())
    print(f"Sheet1 : {len(sheet1_tickers)} tickers existants")

    missing_sheet1 = {t: n for t, n in etf_defensif.items() if t not in sheet1_tickers}
    print(f"Manquants dans Sheet1 : {len(missing_sheet1)}")

    tickers_csv_existing = _load_tickers_csv(tk_path)
    missing_csv = sorted(t for t in etf_defensif if t not in tickers_csv_existing)
    print(f"Manquants dans tickers.csv : {len(missing_csv)}")

    if not write:
        for t, n in list(missing_sheet1.items())[:10]:
            s2, s3, s4, s5 = classify(n)
            print(f"  [dry] {t:12} {n[:50]:50} -> {s2} > {s3} > {s4} > {s5}")
        print("[dry] Rien écrit (--dry).")
        return

    ts = dt.datetime.now().strftime("%Y%m%d-%H%M%S")

    # 1) Sheet1 : une ligne par ticker manquant, Secteur 1 = ETF + classification.
    next_row = ws.max_row + 1
    n_written = 0
    for t, n in sorted(missing_sheet1.items()):
        s2, s3, s4, s5 = classify(n)
        ws.cell(row=next_row, column=col["Ticker Yahoo Finance"], value=t)
        ws.cell(row=next_row, column=col["Nom"], value=n)
        ws.cell(row=next_row, column=col["Secteur"], value="ETF")
        ws.cell(row=next_row, column=col["Secteur 1"], value="ETF")
        ws.cell(row=next_row, column=col["Secteur 2"], value=s2)
        ws.cell(row=next_row, column=col["Secteur 3"], value=s3)
        ws.cell(row=next_row, column=col["Secteur 4"], value=s4)
        ws.cell(row=next_row, column=col["Secteur 5"], value=s5)
        next_row += 1
        n_written += 1

    bak_xlsx = backup_file(
        bf,
        category="maintenance/reinject_missing_etfs",
        filename=f"{bf.stem}.bak-{ts}{bf.suffix}",
    )
    bak_csv = None
    if missing_csv:
        bak_csv = backup_file(
            tk_path,
            category="maintenance/reinject_missing_etfs",
            filename=f"{tk_path.stem}.bak-{ts}{tk_path.suffix}",
        )
    wb.save(bf)
    print(f"[OK] {n_written} lignes ajoutées à Sheet1 (backup : {bak_xlsx.name})")

    # 2) tickers.csv : ajoute les tickers absents (Ticker;Nom;Bourse;Type).
    if missing_csv:
        with open(tk_path, "a", encoding="utf-8", newline="") as f:
            writer = csv.writer(f, delimiter=";")
            for t in missing_csv:
                writer.writerow([t, etf_defensif[t], "ETF USA", "Tracker/ETF"])
        print(f"[OK] {len(missing_csv)} tickers ajoutés à tickers.csv (backup : {bak_csv.name})")

    try:
        from app.services.finance.buffett.broker_availability import reset_etf_cache
        reset_etf_cache()
    except Exception:
        pass


if __name__ == "__main__":
    main(write="--dry" not in sys.argv)

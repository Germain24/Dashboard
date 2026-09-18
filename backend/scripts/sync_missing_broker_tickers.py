"""Ajoute à ToutBroker les tickers du catalogue qui n'y ont aucune ligne.

Les nouveaux instruments sont marqués indisponibles chez les brokers déjà
inventoriés (Trading 212 et les deux comptes Bourse Direct). IBKR reste vide :
son catalogue n'a pas encore été importé.

Usage depuis ``backend/``::

    .venv/Scripts/python.exe scripts/sync_missing_broker_tickers.py --dry-run
    .venv/Scripts/python.exe scripts/sync_missing_broker_tickers.py --apply
"""

from __future__ import annotations

import argparse
import os
import re
import sys
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape
from zipfile import ZIP_DEFLATED, ZipFile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook  # noqa: E402

from app.services.backup_storage import backup_file  # noqa: E402
from app.services.finance.buffett.config import Config  # noqa: E402
from app.services.finance.buffett.ticker_universe import read_ticker_catalog  # noqa: E402

DEFAULT_TICKERS = ROOT.parent / "data/imports/Finances/variables/tickers.csv"
TICKER_COLUMN = "Ticker Yahoo Finance"
ZERO_COLUMNS = ("Tradding 212", "Bourse Direct", "Bourse Direct 2")
IBKR_COLUMN = "IBKR"


def _normalized(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip().upper()


def _inline_cell(column: str, row: int, value: str) -> str:
    safe = escape(value).replace("\r", "&#13;")
    return f'<c r="{column}{row}" t="inlineStr"><is><t>{safe}</t></is></c>'


def _append_rows_to_xlsx(
    source: Path,
    destination: Path,
    missing: list[str],
    rows_by_ticker: dict[str, tuple[str, ...]],
    first_new_row: int,
) -> None:
    """Copie le XLSX octet pour octet sauf sheet1.xml, où les lignes sont ajoutées.

    Cette voie conserve styles, largeurs, formules, filtres et feuilles annexes.
    Une reconstruction OpenPyXL normale matérialiserait plus de cinq millions de
    cellules et prendrait des dizaines de minutes.
    """
    with ZipFile(source, "r") as archive_in, ZipFile(
        destination, "w", compression=ZIP_DEFLATED, allowZip64=True
    ) as archive_out:
        for info in archive_in.infolist():
            payload = archive_in.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                final_row = first_new_row + len(missing) - 1
                payload = re.sub(
                    rb'<dimension ref="[^"]*"\s*/>',
                    f'<dimension ref="A1:Z{final_row}" />'.encode(),
                    payload,
                    count=1,
                )
                marker = b"</sheetData>"
                position = payload.find(marker)
                if position < 0:
                    raise RuntimeError("sheetData introuvable dans la feuille principale")
                with archive_out.open(info, "w", force_zip64=True) as target:
                    target.write(payload[:position])
                    for offset, ticker in enumerate(missing):
                        row_number = first_new_row + offset
                        source_row = rows_by_ticker[ticker]
                        cells = [_inline_cell("A", row_number, ticker)]
                        if len(source_row) > 1 and source_row[1]:
                            cells.append(_inline_cell("B", row_number, source_row[1]))
                        cells.extend(
                            f'<c r="{column}{row_number}" t="n"><v>0</v></c>'
                            for column in ("O", "P", "Q")
                        )
                        if len(source_row) > 3 and source_row[3].strip().upper() == "ETF":
                            cells.append(_inline_cell("S", row_number, "ETF"))
                        target.write(
                            f'<row r="{row_number}">{"".join(cells)}</row>'.encode()
                        )
                    target.write(payload[position:])
                continue
            archive_out.writestr(info, payload)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--dry-run", action="store_true")
    mode.add_argument("--apply", action="store_true")
    parser.add_argument("--tickers", type=Path, default=DEFAULT_TICKERS)
    parser.add_argument("--broker", type=Path, default=Path(Config.BROKER_FILE))
    parser.add_argument("--output", type=Path)
    args = parser.parse_args(argv)

    if not args.tickers.exists() or not args.broker.exists():
        print(f"fichier absent : tickers={args.tickers}, broker={args.broker}")
        return 1

    catalog = read_ticker_catalog(args.tickers, require_canonical=True)
    first_row_by_ticker: dict[str, tuple[str, ...]] = {}
    for row in catalog.rows:
        first_row_by_ticker.setdefault(row[0], row)

    broker_handle = args.broker.open("rb")
    workbook = load_workbook(broker_handle, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    headers = {
        str(cell.value).strip(): index
        for index, cell in enumerate(sheet[1], start=1)
        if cell.value is not None
    }
    required = (TICKER_COLUMN, *ZERO_COLUMNS, IBKR_COLUMN)
    absent = [column for column in required if column not in headers]
    if absent:
        print(f"colonnes absentes de {sheet.title!r} : {absent}")
        workbook.close()
        broker_handle.close()
        return 1

    existing = {
        _normalized(row[0])
        for row in sheet.iter_rows(
            min_row=2,
            min_col=headers[TICKER_COLUMN],
            max_col=headers[TICKER_COLUMN],
            values_only=True,
        )
        if _normalized(row[0])
    }
    missing = [ticker for ticker in catalog.tickers if ticker not in existing]
    print(f"catalogue : {len(catalog.tickers):,} tickers")
    print(f"ToutBroker : {len(existing):,} tickers uniques")
    print(f"à ajouter : {len(missing):,}")
    print("nouveaux brokers : Tradding 212=0, Bourse Direct=0, "
          "Bourse Direct 2=0, IBKR=vide")

    if args.dry_run or not missing:
        workbook.close()
        broker_handle.close()
        print("--dry-run : aucun fichier écrit." if args.dry_run else "Déjà synchronisé.")
        return 0

    workbook.close()
    broker_handle.close()
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    destination = args.output or args.broker
    backup = None
    if destination.exists():
        backup = backup_file(
            destination,
            category="maintenance/sync_missing_broker_tickers",
            filename=f"{destination.name}.bak-sync-{timestamp}",
        )
    temporary = destination.with_name(f".{destination.name}.{timestamp}.tmp.xlsx")
    try:
        _append_rows_to_xlsx(
            args.broker,
            temporary,
            missing,
            first_row_by_ticker,
            sheet.max_row + 1,
        )
        os.replace(temporary, destination)
    finally:
        if temporary.exists():
            temporary.unlink()

    if backup is not None:
        print(f"sauvegarde : {backup}")
    print(f"écrit : {destination} ({len(missing):,} lignes ajoutées)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

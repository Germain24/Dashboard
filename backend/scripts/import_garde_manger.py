"""Safe, idempotent seed of the personal pantry from ``Garde Mange.xlsx``.

The source workbook is only opened read-only. The default invocation is a dry
run; pass ``--apply`` to write the dashboard pantry and its import manifest.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import unicodedata
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO_ROOT / "backend"))
DEFAULT_SOURCE = Path.home() / "Downloads" / "Garde Mange.xlsx"
DEFAULT_PANTRY = REPO_ROOT / "data" / "cuisine_pantry.json"
DEFAULT_MANIFEST = REPO_ROOT / "data" / "cuisine_pantry_imports.json"
IMPORT_NAME = "garde-manger-xlsx-v1"


def _normalise(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return " ".join(text.casefold().strip().split())


_SOURCE_URLS = {
    _normalise("Multi Vitamin and mineral advanced"): (
        "https://naturalfoodbarn.com/product/multi-vitamin-mineral-advanced-90-tabs/"
    ),
}

_NUTRITION_LABELS: dict[str, dict[str, Any]] = {
    _normalise("Huile d'olive extravierge"): {
        "serving_size": "10 mL",
        "values": [
            {"nutrient": "Énergie", "value": 80, "unit": "kcal"},
            {"nutrient": "Lipides", "value": 9, "unit": "g"},
            {"nutrient": "Acides gras saturés", "value": 1.5, "unit": "g"},
            {"nutrient": "Glucides", "value": 0, "unit": "g"},
            {"nutrient": "Protéines", "value": 0, "unit": "g"},
            {"nutrient": "Sodium", "value": 0, "unit": "mg"},
        ],
        "source_label": "Kirkland Signature Organic Extra Virgin Olive Oil 2 L",
        "source_url": "https://www.costco.ca/p/-/kirkland-signature-organic-extra-virgin-olive-oil-2-l/100416825?langId=-25",
    },
    _normalise("Avoine"): {
        "serving_size": "40 g",
        "values": [
            {"nutrient": "Énergie", "value": 150, "unit": "kcal"},
            {"nutrient": "Lipides", "value": 3, "unit": "g"},
            {"nutrient": "Acides gras saturés", "value": 0.5, "unit": "g"},
            {"nutrient": "Glucides", "value": 27, "unit": "g"},
            {"nutrient": "Fibres", "value": 4, "unit": "g"},
            {"nutrient": "Sucres", "value": 1, "unit": "g"},
            {"nutrient": "Protéines", "value": 5, "unit": "g"},
            {"nutrient": "Vitamine D", "value": 0, "unit": "µg"},
            {"nutrient": "Calcium", "value": 20, "unit": "mg"},
            {"nutrient": "Fer", "value": 1.5, "unit": "mg"},
            {"nutrient": "Potassium", "value": 150, "unit": "mg"},
            {"nutrient": "Thiamine (B1)", "value": 0.2, "unit": "mg"},
            {"nutrient": "Phosphore", "value": 130, "unit": "mg"},
            {"nutrient": "Magnésium", "value": 40, "unit": "mg"},
        ],
        "source_label": "Quaker Old Fashioned Oats (UPC 030000010204)",
        "source_url": "https://smartlabel.pepsico.info/030000010204-0001-en-US/index.html",
        "source_updated": "2022-05-24",
        "note": "Étiquette Quaker Old Fashioned; confirmer la variété exacte du sac de 10,26 kg.",
    },
    _normalise("Graine de chia"): {
        "serving_size": "15 g",
        "values": [
            {"nutrient": "Énergie", "value": 70, "unit": "kcal"},
            {"nutrient": "Lipides", "value": 4.5, "unit": "g"},
            {"nutrient": "Acides gras saturés", "value": 0.5, "unit": "g"},
            {"nutrient": "Sodium", "value": 2, "unit": "mg"},
            {"nutrient": "Glucides", "value": 6, "unit": "g"},
            {"nutrient": "Fibres", "value": 5, "unit": "g"},
            {"nutrient": "Sucres", "value": 0, "unit": "g"},
            {"nutrient": "Protéines", "value": 2, "unit": "g"},
            {"nutrient": "Calcium", "value": 100, "unit": "mg"},
            {"nutrient": "Fer", "value": 1.25, "unit": "mg"},
            {"nutrient": "Potassium", "value": 50, "unit": "mg"},
        ],
        "source_label": "Kirkland Signature Organic Chia Seeds, 1.36 kg",
        "source_url": "https://www.costco.ca/p/-/kirkland-signature-organic-chia-seeds-136-kg/100539558?langId=-24",
    },
    _normalise("Spaghettini"): {
        "serving_size": "85 g (sec)",
        "values": [
            {"nutrient": "Énergie", "value": 300, "unit": "kcal"},
            {"nutrient": "Lipides", "value": 1.5, "unit": "g"},
            {"nutrient": "Acides gras saturés", "value": 0.4, "unit": "g"},
            {"nutrient": "Glucides", "value": 63, "unit": "g"},
            {"nutrient": "Fibres", "value": 5, "unit": "g"},
            {"nutrient": "Sucres", "value": 2, "unit": "g"},
            {"nutrient": "Protéines", "value": 11, "unit": "g"},
            {"nutrient": "Cholestérol", "value": 0, "unit": "mg"},
            {"nutrient": "Calcium", "value": 20, "unit": "mg"},
            {"nutrient": "Fer", "value": 2.5, "unit": "mg"},
        ],
        "source_label": "Barilla Canada Spaghettini",
        "source_url": "https://www.barilla.com/fr-ca/produits/pates/bo%C3%AEte-bleue-classique/spaghettini",
        "note": "Le sodium affiché sur la page est incohérent (3 g / 1 % VQ); valeur écartée.",
    },
}


def _source_columns(headers: list[Any]) -> tuple[int, int, int]:
    normalized = [_normalise(value) for value in headers]
    product = next((i for i, value in enumerate(normalized) if value == "produit"), None)
    brand = next(
        (i for i, value in enumerate(normalized) if value in {"marque", "martque"}),
        None,
    )
    quantity = next((i for i, value in enumerate(normalized) if value == "quantite"), None)
    if product is None or brand is None or quantity is None:
        raise ValueError("Le classeur doit contenir Produit, Marque/Martque et Quantité.")
    return product, brand, quantity


def _parse_quantity(raw: Any) -> tuple[float, str]:
    text = str(raw or "").strip()
    match = re.fullmatch(r"([+-]?\d+(?:[.,]\d+)?)\s*([^\d\s].*?)?", text)
    if not match:
        raise ValueError(f"Quantité invalide dans le classeur : {text!r}")
    quantity = float(match.group(1).replace(",", "."))
    unit_raw = (match.group(2) or "unité").strip()
    unit_key = _normalise(unit_raw).replace(".", "")
    known_units = {
        "g": "g", "gramme": "g", "grammes": "g",
        "kg": "kg", "kilogramme": "kg", "kilogrammes": "kg",
        "mg": "mg", "milligramme": "mg", "milligrammes": "mg",
        "ml": "mL", "millilitre": "mL", "millilitres": "mL",
        "l": "L", "litre": "L", "litres": "L",
        "tablet": "tablets", "tablets": "tablets", "tablette": "tablets",
        "tablettes": "tablets", "comprime": "comprimés", "comprimes": "comprimés",
    }
    return quantity, known_units.get(unit_key, unit_raw)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


# Explicit, exact source-name mappings only. Names and codes were checked
# against data/imports/Sante/ciqual/ciqual_normalise.csv and the current food
# optimizer catalog. Deliberately ambiguous packaged foods remain unmatched.
_MATCHES: dict[str, dict[str, str]] = {
    _normalise("Huile d'olive extravierge"): {
        "ingredient": "Huile d'olive extra-vierge",
        "ciqual_code": "17270",
        "ciqual_nom": "Huile d'olive vierge extra",
    },
    _normalise("Avoine"): {
        "ingredient": "Flocons d'avoine",
        "ciqual_code": "9311",
        "ciqual_nom": "Flocon d'avoine",
    },
    _normalise("Graine de chia"): {
        "ingredient": "Graines de chia",
        "ciqual_code": "15047",
        "ciqual_nom": "Chia, graine, séchée",
    },
    _normalise("Riz basmati"): {
        "ingredient": "Riz basmati (sec)",
        "ciqual_code": "9119",
        "ciqual_nom": "Riz thaï ou basmati, cru",
    },
    _normalise("Spaghettini"): {
        "ingredient": "Pates (sec)",
        "ciqual_code": "9810",
        "ciqual_nom": "Pâtes sèches standard, crues",
    },
    _normalise("Vinaigre balsamique"): {
        "ciqual_code": "11091",
        "ciqual_nom": "Vinaigre balsamique",
    },
    _normalise("Haricots verts en conserve"): {
        "ciqual_code": "20062",
        "ciqual_nom": "Haricot vert, appertisé, égoutté",
    },
}

_SUPPLEMENTS = {
    _normalise(name)
    for name in (
        "Pre-Workout Anabolix Freaked",
        "L-Glutamine",
        "Amino Acids Dragon Fuel",
        "Creatne Monohydrate",
        "Custard Casein Protein",
        "Beta Alanine",
        "Multi Vitamin and mineral advanced",
        "Whey Protein Vanille",
    )
}


def read_workbook(source: Path) -> tuple[list[dict[str, Any]], str, str]:
    """Read workbook rows without ever modifying the input file."""
    source_hash = _sha256(source)
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = list(next(rows, ()))
        product_column, brand_column, quantity_column = _source_columns(headers)
        parsed: list[dict[str, Any]] = []
        for row_number, row in enumerate(rows, start=2):
            values = list(row)
            product = str(values[product_column] or "").strip() if product_column < len(values) else ""
            brand = str(values[brand_column] or "").strip() if brand_column < len(values) else ""
            raw_quantity = values[quantity_column] if quantity_column < len(values) else None
            if not product and not brand and raw_quantity is None:
                continue
            if not product or raw_quantity is None:
                raise ValueError(f"Ligne {row_number} incomplète dans le classeur.")
            quantity, unit = _parse_quantity(raw_quantity)
            if quantity <= 0:
                raise ValueError(f"Quantité non positive à la ligne {row_number}.")

            parsed.append({
                "source_row": row_number,
                "source_produit": product,
                "marque": brand,
                "quantite": quantity,
                "unite": unit,
                "quantite_source": str(raw_quantity).strip(),
                "sheet": sheet.title,
            })
        return parsed, source_hash, sheet.title
    finally:
        workbook.close()


def _import_key(row: dict[str, Any], occurrence: int) -> str:
    identity = f"{_normalise(row['source_produit'])}\0{_normalise(row['marque'])}\0{occurrence}"
    return hashlib.sha256(identity.encode("utf-8")).hexdigest()


def _pantry_row(
    row: dict[str, Any], source_hash: str, occurrence: int, source_name: str
) -> dict[str, Any]:
    match = _MATCHES.get(_normalise(row["source_produit"]), {})
    is_supplement = _normalise(row["source_produit"]) in _SUPPLEMENTS
    ingredient = match.get("ingredient") or row["source_produit"]
    if is_supplement:
        aisle = "Compléments"
    elif "conserve" in _normalise(row["source_produit"]):
        aisle = "Conserves"
    elif row["unite"] in {"mL", "L"}:
        aisle = "Épicerie sèche" if "huile" in _normalise(row["source_produit"]) or "vinaigre" in _normalise(row["source_produit"]) else "Conserves"
    else:
        aisle = "Épicerie sèche"

    item: dict[str, Any] = {
        "ingredient": ingredient,
        "quantite": row["quantite"],
        "unite": row["unite"],
        "date_peremption": None,
        "rayon": aisle,
        "source_produit": row["source_produit"],
        "marque": row["marque"],
        "quantite_source": row["quantite_source"],
        "type_aliment": "supplement" if is_supplement else ("aliment" if match.get("ingredient") else "non_mappe"),
        "source_import": IMPORT_NAME,
        "source_workbook": source_name,
        "source_sheet": row["sheet"],
        "source_row": row["source_row"],
        "source_sha256": source_hash,
        "source_import_key": _import_key(row, occurrence),
    }
    if is_supplement:
        item["optimizer_excluded"] = True
    source_url = _SOURCE_URLS.get(_normalise(row["source_produit"]))
    if source_url:
        item["source_url"] = source_url
    nutrition_label = _NUTRITION_LABELS.get(_normalise(row["source_produit"]))
    if nutrition_label:
        item["nutrition_label"] = nutrition_label
    if match.get("ciqual_code"):
        item["ciqual_code"] = match["ciqual_code"]
        item["ciqual_nom"] = match["ciqual_nom"]
    return item


def _read_json(path: Path, default: Any) -> Any:
    if not path.exists():
        return default
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ValueError(f"Fichier JSON invalide, aucune écriture effectuée : {path}") from exc


def _atomic_write(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    temporary.replace(path)


def _backup_persistent_file(path: Path) -> None:
    """Archive existing default pantry state to the configured NAS before replacement."""
    default_files = {DEFAULT_PANTRY.resolve(), DEFAULT_MANIFEST.resolve()}
    if path.resolve() not in default_files or not path.is_file():
        return
    from app.services.backup_storage import backup_file

    digest = _sha256(path)[:12]
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    backup_file(
        path,
        category="cuisine/pantry",
        filename=f"{path.stem}-{stamp}-{digest}{path.suffix}",
    )


def import_workbook(
    source: Path,
    pantry_path: Path,
    manifest_path: Path,
    *,
    apply: bool = False,
) -> dict[str, Any]:
    """Plan and optionally apply an idempotent import, preserving existing items."""
    source = Path(source)
    pantry_path = Path(pantry_path)
    manifest_path = Path(manifest_path)
    if not source.is_file():
        raise FileNotFoundError(source)

    rows, source_hash, sheet = read_workbook(source)
    pantry = _read_json(pantry_path, [])
    manifest = _read_json(manifest_path, {"imports": {}})
    if not isinstance(pantry, list) or any(not isinstance(item, dict) for item in pantry):
        raise ValueError(f"Le garde-manger n'est pas une liste valide : {pantry_path}")
    if not isinstance(manifest, dict) or not isinstance(manifest.get("imports", {}), dict):
        raise ValueError(f"Le manifeste d'import n'est pas valide : {manifest_path}")

    prior_import = manifest.get("imports", {}).get(IMPORT_NAME)
    same_source = isinstance(prior_import, dict) and prior_import.get("sha256") == source_hash

    existing_by_key = {
        item.get("source_import_key"): index
        for index, item in enumerate(pantry)
        if item.get("source_import_key")
    }
    existing_keys = set(existing_by_key)
    seen_keys: set[str] = set()
    occurrences: Counter[tuple[str, str]] = Counter()
    additions: list[dict[str, Any]] = []
    updates: dict[int, dict[str, Any]] = {}
    for row in rows:
        identity = (_normalise(row["source_produit"]), _normalise(row["marque"]))
        occurrences[identity] += 1
        item = _pantry_row(row, source_hash, occurrences[identity], source.name)
        key = item["source_import_key"]
        if key in existing_keys:
            index = existing_by_key[key]
            existing_item = updates.get(index, pantry[index])
            enriched = dict(existing_item)
            for field in ("nutrition_label",):
                if field not in enriched and field in item:
                    enriched[field] = item[field]
            if enriched != existing_item:
                updates[index] = enriched
            seen_keys.add(key)
            continue
        if key in seen_keys:
            seen_keys.add(key)
            continue
        seen_keys.add(key)
        additions.append(item)

    existing_ids = [item.get("id") for item in pantry if isinstance(item.get("id"), int)]
    next_id = max(existing_ids, default=0) + 1
    for item in additions:
        item["id"] = next_id
        next_id += 1

    result = {
        "applied": apply,
        "already_imported": False,
        "source_sha256": source_hash,
        "source_sheet": sheet,
        "rows": len(rows),
        "added": len(additions),
        "updated": len(updates),
        "skipped": len(rows) - len(additions),
        "already_imported": same_source and not additions and not updates,
        "supplements": sum(item["type_aliment"] == "supplement" for item in additions),
        "ciqual_matched": sum("ciqual_code" in item for item in additions),
        "pantry_path": str(pantry_path),
    }
    if not apply or (same_source and not additions and not updates):
        return result

    manifest.setdefault("imports", {})[IMPORT_NAME] = {
        "sha256": source_hash,
        "source_workbook": source.name,
        "source_sheet": sheet,
        "source_rows": len(rows),
        "added_rows": len(additions),
        "completed_at": datetime.now(timezone.utc).isoformat(),
    }
    # The row-level keys make a retry safe if the process stops between writes.
    if additions or updates:
        _backup_persistent_file(pantry_path)
        _atomic_write(pantry_path, [
            updates.get(index, item) for index, item in enumerate(pantry)
        ] + additions)
    if additions or not same_source:
        _backup_persistent_file(manifest_path)
        _atomic_write(manifest_path, manifest)
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--pantry", type=Path, default=DEFAULT_PANTRY)
    parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    parser.add_argument("--apply", action="store_true", help="Écrit le garde-manger; défaut : aperçu seulement.")
    args = parser.parse_args()
    print(json.dumps(
        import_workbook(args.source, args.pantry, args.manifest, apply=args.apply),
        ensure_ascii=False,
        indent=2,
    ))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

"""Génère la feuille "Prix" de `data/imports/Voyage.xlsx` : matrice de prix
estimés (formule par distance, cf. `app.services.voyage.price_estimator`),
aéroports de DÉPART en colonnes (X) et aéroports d'ARRIVÉE en lignes (Y),
pour chaque aéroport de la wishlist Voyage + YUL -- éditable à la main,
prioritaire sur la formule au runtime (relu à chaque planification).

Script one-shot, pas couvert par les tests automatisés (lit la DB, écrit
Voyage.xlsx) : à relancer manuellement si la wishlist Voyage change
significativement (nouvel import Excel avec de nouveaux aéroports). Un
backup horodaté de Voyage.xlsx est fait avant toute écriture.

Usage (depuis `backend/`) : `uv run python -m scripts.build_prix_vols_estimes`
"""
from __future__ import annotations

import datetime as dt
from itertools import combinations

import openpyxl
from sqlmodel import Session, select

from app.core.config import settings
from app.core.db import engine
from app.models.voyage import LieuVoyage
from app.services.backup_storage import backup_file
from app.services.voyage.airports import haversine_km, lookup_coords
from app.services.voyage.price_estimator import SHEET_PRIX, estimate_price_eur

_HOME_AIRPORTS = ["YUL"]  # aéroports de départ habituels, en plus des destinations


def main() -> None:
    with Session(engine) as s:
        lieux = s.exec(select(LieuVoyage)).all()
        iatas = sorted({lv.aeroport_iata for lv in lieux if lv.aeroport_iata} | set(_HOME_AIRPORTS))

    prix: dict[frozenset[str], float] = {}
    for a, b in combinations(iatas, 2):
        coords_a, coords_b = lookup_coords(a), lookup_coords(b)
        if coords_a is None or coords_b is None:
            continue
        p, _duree_min = estimate_price_eur(haversine_km(coords_a, coords_b))
        prix[frozenset((a, b))] = p

    path = settings.imports_dir / "Voyage.xlsx"
    ts = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = backup_file(
        path,
        category="maintenance/build_prix_vols_estimes",
        filename=f"{path.stem}.backup-{ts}{path.suffix}",
    )

    wb = openpyxl.load_workbook(path)
    if SHEET_PRIX in wb.sheetnames:
        del wb[SHEET_PRIX]
    ws = wb.create_sheet(SHEET_PRIX)

    ws.append(["Arrivée \\ Départ", *iatas])
    for arrivee in iatas:
        row = [arrivee]
        for depart in iatas:
            if depart == arrivee:
                row.append(None)
            else:
                row.append(prix.get(frozenset((depart, arrivee))))
        ws.append(row)

    wb.save(path)
    print(f"feuille '{SHEET_PRIX}' écrite ({len(iatas)} aéroports) -> {path} (backup : {backup.name})")


if __name__ == "__main__":
    main()

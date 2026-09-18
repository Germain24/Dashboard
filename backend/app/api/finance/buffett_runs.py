"""Routes Buffett : gestion des runs, progression, exports."""

from __future__ import annotations

import io
import logging
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlmodel import Session, select

from app.api.schemas_finance import (
    BuffettProgressOut,
    BuffettRunDetailOut,
    BuffettRunOut,
)
from app.core.db import get_session
from app.models.finance import BuffettRun, BuffettRunResult, BuffettRunStatus

logger = logging.getLogger(__name__)
router = APIRouter()


def _top_score_results(session: Session, run_id: int) -> list[BuffettRunResult]:
    from app.services.finance.buffett.reporting import top_company_results

    return top_company_results(session, run_id)


@router.get("/buffett/runs", response_model=list[BuffettRunOut])
def buffett_runs(limit: int = 12, session: Session = Depends(get_session)):
    stmt = select(BuffettRun).order_by(BuffettRun.run_date.desc()).limit(limit)
    return list(session.exec(stmt).all())


@router.get("/buffett/latest", response_model=Optional[BuffettRunOut])
def buffett_latest(session: Session = Depends(get_session)):
    stmt = (
        select(BuffettRun)
        .where(BuffettRun.statut == BuffettRunStatus.TERMINE.value)
        .order_by(BuffettRun.run_date.desc())
    )
    return session.exec(stmt).first()


@router.get("/buffett/progress", response_model=BuffettProgressOut)
def buffett_progress(session: Session = Depends(get_session)):
    from app.services.finance.scheduler_stub import is_analysis_running

    active = is_analysis_running()

    if not active:
        stuck = session.exec(
            select(BuffettRun).where(BuffettRun.statut == BuffettRunStatus.EN_COURS.value)  # type: ignore[attr-defined]
        ).all()
        changed = False
        for sr in stuck:
            sr.statut = BuffettRunStatus.INTERROMPU.value
            sr.erreur = "Process interrompu (relancez pour reprendre)"
            session.add(sr)
            changed = True
        if changed:
            session.commit()

    run = session.exec(
        select(BuffettRun)
        .where(
            BuffettRun.statut.in_([BuffettRunStatus.EN_COURS.value, BuffettRunStatus.INTERROMPU.value])
        )  # type: ignore[attr-defined]
        .order_by(BuffettRun.created_at.desc())
    ).first()
    if run:
        from app.services.finance.buffett.progress_state import snapshot as progress_snapshot
        from app.services.finance.buffett.rate_limiter import active_paused_until
        extra = progress_snapshot()

        extra.update(
            run_id=run.id,
            statut=run.statut,
            active=active,
            progress_pct=run.progress_pct or 0.0,
            n_done=run.n_tickers_analyzed,
            n_total=run.n_tickers_total,
            paused_until=active_paused_until() if active else None,
        )
        return BuffettProgressOut(**extra)
    return BuffettProgressOut(run_id=None, statut="idle", active=False, progress_pct=0.0)


@router.get("/portfolio/progress")
def portfolio_progress(history_after: int | None = Query(default=None, ge=0)):
    from app.services.finance.buffett import optimization_progress as opt_prog

    snap = opt_prog.snapshot(history_after=history_after)
    snap["progress_pct"] = round(snap.get("convergence", 0.0) * 100, 1)
    return snap


@router.post("/buffett/optimization/stop")
def optimization_stop():
    from app.services.finance.buffett import optimization_progress as opt_prog
    from app.services.finance.buffett import progress_state

    accepted = opt_prog.request_stop() or progress_state.request_stop()
    return {
        "accepted": accepted,
        "message": (
            "Arrêt demandé — le meilleur portefeuille sera conservé."
            if accepted else "Aucune optimisation active."
        ),
    }


@router.get("/buffett/runs/{run_id}", response_model=BuffettRunDetailOut)
def buffett_run_detail(run_id: int, session: Session = Depends(get_session)):
    run = session.get(BuffettRun, run_id)
    if not run:
        raise HTTPException(404, f"Run {run_id} introuvable")
    top = _top_score_results(session, run_id)
    alloc = list(
        session.exec(
            select(BuffettRunResult)
            .where(BuffettRunResult.run_id == run_id)
            .where(BuffettRunResult.allocation_pct.isnot(None))
            .order_by(BuffettRunResult.allocation_pct.desc())
        ).all()
    )
    optimization = (run.params_json or {}).get("optimization")
    return BuffettRunDetailOut(
        run=run, top_results=top, allocation_cible=alloc, optimization=optimization,
    )


@router.delete("/buffett/runs/{run_id}", status_code=204)
def buffett_run_delete(run_id: int, session: Session = Depends(get_session)):
    from app.services.finance.buffett.reporting import delete_run

    if not delete_run(session, run_id):
        raise HTTPException(404, f"Run {run_id} introuvable")


@router.get("/buffett/runs/{run_id}/export")
def buffett_run_export(run_id: int, session: Session = Depends(get_session)):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    run = session.get(BuffettRun, run_id)
    if not run:
        raise HTTPException(404, f"Run {run_id} introuvable")

    results = list(
        session.exec(
            select(BuffettRunResult)
            .where(BuffettRunResult.run_id == run_id)
            .order_by(BuffettRunResult.chance_moat.desc())
        ).all()
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resultats MOAT"

    ws.append([f"Analyse Buffett - Run #{run_id} - {run.run_date}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([
        f"Statut : {run.statut}",
        f"Tickers analyses : {run.n_tickers_analyzed or 0}/{run.n_tickers_total or 0}",
        f"Duree : {run.duree_sec or 0:.0f}s",
    ])
    ws.append([])

    headers = ["Ticker", "Nom", "Secteur", "Pays", "Score MOAT", "Alloc. cible (%)", "Broker cible", "Achat"]
    header_row = 4
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="1E3A5F")
    header_font = Font(bold=True, color="FFFFFF")
    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    green_fill = PatternFill("solid", fgColor="D6F5D6")
    blue_fill = PatternFill("solid", fgColor="D6E8F5")
    for r in results:
        score_val = r.chance_moat
        row = [
            r.ticker, r.nom or "", r.secteur or "", r.pays or "",
            round(score_val, 2) if score_val is not None else None,
            round(r.allocation_pct, 2) if r.allocation_pct is not None else None,
            r.broker_cible or "", "oui" if r.achat else "",
        ]
        ws.append(row)
        if (r.secteur or "").strip().upper() == "ETF":
            fill = blue_fill
        elif score_val and score_val >= 80:
            fill = green_fill
        else:
            fill = None
        if fill:
            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=ws.max_row, column=col_idx).fill = fill

    col_widths = [12, 32, 20, 18, 12, 16, 16, 7]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    alloc_list = [r for r in results if r.allocation_pct]
    if alloc_list:
        ws2 = wb.create_sheet("Allocation cible")
        ws2.append(["Ticker", "Nom", "Broker cible", "Alloc. cible (%)"])
        for cell in ws2[1]:
            cell.fill = header_fill
            cell.font = header_font
        for r in sorted(alloc_list, key=lambda x: -(x.allocation_pct or 0)):
            ws2.append([r.ticker, r.nom or "", r.broker_cible or "", round(r.allocation_pct, 2) if r.allocation_pct else None])
        for i, w in enumerate([12, 32, 16, 16], 1):
            ws2.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"buffett_run_{run_id}_{run.run_date}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/buffett/runs/{run_id}/export.csv")
def buffett_run_export_csv(run_id: int, session: Session = Depends(get_session)):
    import csv
    import io as _io

    run = session.get(BuffettRun, run_id)
    if not run:
        raise HTTPException(404, f"Run {run_id} introuvable")

    results = list(
        session.exec(
            select(BuffettRunResult)
            .where(BuffettRunResult.run_id == run_id)
            .order_by(BuffettRunResult.chance_moat.desc())
        ).all()
    )

    buf = _io.StringIO()
    writer = csv.writer(buf, delimiter=";")
    writer.writerow([
        "Ticker", "Nom", "Secteur", "Pays", "Score MOAT", "Achat",
        "Allocation cible (%)", "Broker cible", "PER", "Prix",
    ])
    for r in results:
        writer.writerow([
            r.ticker, r.nom or "", r.secteur or "", r.pays or "",
            round(r.chance_moat, 2) if r.chance_moat is not None else "",
            "oui" if r.achat else "",
            round(r.allocation_pct, 2) if r.allocation_pct is not None else "",
            r.broker_cible or "",
            round(r.per, 2) if r.per else "",
            round(r.prix, 2) if r.prix else "",
        ])
    buf.seek(0)
    filename = f"buffett_run_{run_id}_{run.run_date}.csv"
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )

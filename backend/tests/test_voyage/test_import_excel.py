"""Import Voyage.xlsx -> cache lieu_voyage."""
from __future__ import annotations

import openpyxl
import pytest
from sqlmodel import Session, SQLModel, create_engine, select

import app.models  # noqa: F401
from app.models.voyage import LieuVoyage
from app.services.voyage.import_excel import parse_voyage_xlsx, sync_voyage


@pytest.fixture(name="session")
def session_fixture():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False})
    SQLModel.metadata.create_all(engine)
    with Session(engine) as s:
        yield s


def _make_xlsx(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Lieux", "Ville (ou ville la plus proche)", "Pays", "Visité", "Ordre",
               "Aéroport (IATA)", "Jours min", "Jours max", "Coût/jour estimé"])
    ws.append(["Table Mountain", "Le Cap", "Afrique du Sud", False, 1, "CPT", 2, 4, 80])
    ws.append(["K-2", "K-2", "Chine", False, 5, None, None, None, None])  # incomplet
    ws.append(["Robben Island", "Le Cap", "Afrique du Sud", True, None, "CPT", 1, 1, 60])  # déjà visité
    ws.append([None, None, None, None, None, None, None, None, None])  # ligne vide -> ignorée
    wb.save(path)


def test_parse_voyage_xlsx(tmp_path):
    p = tmp_path / "Voyage.xlsx"
    _make_xlsx(p)
    rows = parse_voyage_xlsx(p)
    assert len(rows) == 3
    assert rows[0] == {
        "nom": "Table Mountain", "ville": "Le Cap", "pays": "Afrique du Sud", "visite": False,
        "aeroport_iata": "CPT", "jours_min": 2, "jours_max": 4, "cout_jour_estime": 80.0,
    }
    assert rows[1]["aeroport_iata"] is None
    assert rows[1]["jours_min"] is None
    assert rows[2]["visite"] is True


def test_sync_voyage_wipes_and_refills_and_flags_incomplete(tmp_path, session):
    p = tmp_path / "Voyage.xlsx"
    _make_xlsx(p)
    session.add(LieuVoyage(nom="Obsolète"))
    session.commit()

    result = sync_voyage(session, p)

    assert result["lieux"] == 3
    assert result["incomplets"] == ["K-2"]  # non visité + aéroport/jours manquants
    noms = {lv.nom for lv in session.exec(select(LieuVoyage)).all()}
    assert noms == {"Table Mountain", "K-2", "Robben Island"}  # ancien "Obsolète" effacé


def test_sync_voyage_flags_incomplete_when_cout_jour_estime_missing(tmp_path, session):
    """Issue 3 (revue finale) : aéroport + jours renseignés mais coût/jour manquant
    -> incomplet (cohérent avec `_est_complet` côté API)."""
    p = tmp_path / "Voyage.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Lieux", "Ville (ou ville la plus proche)", "Pays", "Visité", "Ordre",
               "Aéroport (IATA)", "Jours min", "Jours max", "Coût/jour estimé"])
    ws.append(["Sans coût", "Ville X", "Pays X", False, 1, "CPT", 2, 4, None])
    wb.save(p)

    result = sync_voyage(session, p)
    assert result["incomplets"] == ["Sans coût"]


def test_sync_voyage_does_not_flag_incomplete_if_already_visited(tmp_path, session):
    p = tmp_path / "Voyage.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Lieux", "Ville (ou ville la plus proche)", "Pays", "Visité", "Ordre",
               "Aéroport (IATA)", "Jours min", "Jours max", "Coût/jour estimé"])
    ws.append(["Déjà fait", "Paris", "France", True, None, None, None, None, None])
    wb.save(p)

    result = sync_voyage(session, p)
    assert result["incomplets"] == []


def test_marquer_visites_writes_excel_and_db(tmp_path, session):
    from app.services.voyage.import_excel import marquer_visites

    p = tmp_path / "Voyage.xlsx"
    _make_xlsx(p)
    sync_voyage(session, p)  # peuple la DB depuis l'Excel initial

    backup = marquer_visites(session, p, ["Table Mountain"])

    assert backup.exists()
    assert backup.name.startswith("Voyage.backup-")

    # Excel mis à jour
    import openpyxl
    wb = openpyxl.load_workbook(p, data_only=True)
    ws = wb.active
    header = [c.value for c in ws[1]]
    col_lieux = header.index("Lieux") + 1
    col_visite = header.index("Visité") + 1
    row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, col_lieux).value == "Table Mountain")
    assert ws.cell(row, col_visite).value is True
    # les autres lignes ne sont pas touchées
    row_k2 = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, col_lieux).value == "K-2")
    assert ws.cell(row_k2, col_visite).value is False

    # DB mise à jour dans la même opération
    lv = session.exec(select(LieuVoyage).where(LieuVoyage.nom == "Table Mountain")).first()
    assert lv.visite is True


def test_marquer_visites_raises_if_columns_missing(tmp_path, session):
    from app.services.voyage.import_excel import marquer_visites

    p = tmp_path / "Voyage.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["Autre chose"])
    wb.save(p)

    with pytest.raises(ValueError):
        marquer_visites(session, p, ["X"])


def test_marquer_visites_raises_if_name_not_found(tmp_path, session):
    from app.services.voyage.import_excel import marquer_visites

    p = tmp_path / "Voyage.xlsx"
    _make_xlsx(p)
    sync_voyage(session, p)  # peuple la DB depuis l'Excel initial

    # Try to mark a non-existent name as visited
    with pytest.raises(ValueError, match="Noms introuvables"):
        marquer_visites(session, p, ["Non-existent Place"])

    # Verify Excel was not modified: existing rows should remain unchanged
    wb = openpyxl.load_workbook(p, data_only=True)
    ws = wb.active
    header = [c.value for c in ws[1]]
    col_lieux = header.index("Lieux") + 1
    col_visite = header.index("Visité") + 1

    # Table Mountain should still be False
    row_tm = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, col_lieux).value == "Table Mountain")
    assert ws.cell(row_tm, col_visite).value is False
    wb.close()

    # Verify DB was not modified: Table Mountain should still be visite=False
    lv = session.exec(select(LieuVoyage).where(LieuVoyage.nom == "Table Mountain")).first()
    assert lv.visite is False
